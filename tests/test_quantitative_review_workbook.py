from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import load_workbook
from conftest import fixture_source
from src.schemas.source import Source, SourceState

from src.tools.quantitative_review_workbook import (
    QuantitativeReviewWorkbookError,
    build_quantitative_review_workbook,
    resolve_standard_workbook_template,
)


def _template() -> Path:
    try:
        return resolve_standard_workbook_template()
    except FileNotFoundError:
        pytest.skip("Workbook standar pengguna tidak tersedia di environment test")


def assessed_records(count: int, root: Path, template: Path):
    """Synthetic scoring fixture, using explicit weights only on a copy."""
    wb = load_workbook(template)
    pillars = [wb.worksheets[1].cell(row, 1).value for row in range(5, 8)]
    for row in range(5, 8):
        wb.worksheets[1].cell(row, 2).value = 10
        for col in range(3, 8):
            wb.worksheets[1].cell(row, col).value = f"Synthetic test scoring guide {row}-{col}"
    for row in range(12, 16):
        wb.worksheets[1].cell(row, 3).value = "Synthetic eligible category"
    weighted = root / "synthetic_weighted_template.xlsx"
    wb.save(weighted)
    records = []
    for index in range(1, count + 1):
        source = fixture_source(Source(id=f"src_fixture_{index}", title=f"Synthetic article {index}",
            doi=f"10.1234/fixture.{index}", state=SourceState.APPROVED,
            rights_status="OPEN_LICENSE", access_mode="OPEN_DOWNLOAD", license_url="https://creativecommons.org/licenses/by/4.0/"),
            "Synthetic methods evidence.", root / "sources")
        digest = source.metadata["retrieval"]["sha256"]
        excerpts = [("Methods", "Synthetic methods evidence."), ("Results", "Fixture result."), ("Discussion", "Fixture discussion.")]
        source.metadata["examination"] = {"source_id": source.id, "artifact_sha256": digest,
            "reviewer": "synthetic test reviewer", "reviewed_at": "2026-10-09", "scope": "full",
            "sections": [{"locator": loc, "excerpt": excerpt} for loc, excerpt in excerpts]}
        source.metadata["version"] = "version_of_record"
        source.metadata["eligibility_review"] = {"source_id": source.id, "artifact_sha256": digest,
            "reviewer": "synthetic test reviewer", "reviewed_at": "2026-10-09", "decision": "eligible",
            "sections": [{"criterion": criterion, "locator": loc, "excerpt": excerpt} for criterion, (loc, excerpt)
                in zip(("construct", "population", "design"), excerpts)]}
        # Different verified scores demonstrate sorting rather than input order.
        score = index % 10
        records.append({"title": source.title, "citation": f"Fixture{index} (2025)", "eligible": True,
            "doi": source.doi, "source_snapshot": source.model_dump(mode="json"), "total_score": score * 3,
            "assessment": {"artifact_sha256": digest, "reviewer": "synthetic test reviewer", "reviewed_at": "2026-10-09",
                "criteria": [{"pillar": pillar, "score": score, "locator": loc, "evidence_excerpt": excerpt}
                    for pillar, (loc, excerpt) in zip(pillars, excerpts)]}})
    return records, weighted


@pytest.mark.parametrize("count", [39, 40, 41, 75])
def test_all_records_and_ranks_are_preserved_above_40(tmp_path: Path, count: int) -> None:
    records, template = assessed_records(count, tmp_path, _template())
    result = build_quantitative_review_workbook(records, tmp_path / "review.xlsx", template_path=template, require_free_full_text=True)
    workbook = load_workbook(result.workbook_path)
    master = workbook["Master 40 Kuantitatif PoPCites"]
    ranking = workbook["Rekapitulasi & Peringkat 40"]
    assert [master.cell(6 + i, 1).value for i in range(count)] == list(range(1, count + 1))
    assert [ranking.cell(5 + i, 1).value for i in range(count)] == list(range(1, count + 1))
    assert result.article_count == count
    assert result.ranked_count == count
    assert result.status == "PASS"
    assert result.counts["eligible"] == count
    assert len({ranking.cell(5 + i, 2).value for i in range(count)}) == count
    assert [ranking.cell(5+i, 8).value for i in range(count)] == sorted([r["total_score"] for r in records], reverse=True)
    if count > 40:
        assert master.cell(46, 2).value == "Synthetic article 41"
        assert ranking.cell(45, 1).value == 41
        assert master.cell(46, 2)._style == master.cell(6, 2)._style
    workbook.close()


def test_exact_40_is_explicit_selection(tmp_path: Path) -> None:
    records, template = assessed_records(41, tmp_path, _template())
    result = build_quantitative_review_workbook(records, tmp_path / "review.xlsx", template_path=template, exact_count=40)
    assert result.article_count == 41  # Full corpus retained; final selection is explicit.
    assert result.ranked_count == 40
    assert result.unranked_count == 1


def test_best_without_count_requests_clarification(tmp_path: Path) -> None:
    records, template = assessed_records(4, tmp_path, _template())
    result = build_quantitative_review_workbook(records, tmp_path / "review.xlsx", template_path=template, request_text="artikel terbaik")
    assert result.counts["clarification"]["options"] == [10, 20, 40]
    assert result.article_count == 4 and result.ranked_count == 0 and result.status == "PARTIAL"


def test_doi_without_full_text_is_not_full_text_verified(tmp_path: Path) -> None:
    result = build_quantitative_review_workbook(
        [{"title": "Artikel", "doi": "10.1234/x", "total_score": 1}],
        tmp_path / "review.xlsx",
        template_path=_template(),
    )
    workbook = load_workbook(result.workbook_path)
    status = workbook["Master 40 Kuantitatif PoPCites"].cell(6, 17).value
    assert status == "belum diverifikasi"
