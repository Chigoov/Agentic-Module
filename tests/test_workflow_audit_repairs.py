"""Behavioral regression checks using explicitly synthetic, local artifacts."""
from __future__ import annotations

import argparse
import hashlib
import io
import json
from pathlib import Path
import zipfile

import pytest
from openpyxl import load_workbook
from openpyxl.comments import Comment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table
from openpyxl.workbook.defined_name import DefinedName

from conftest import fixture_source
from test_quantitative_review_workbook import assessed_records, _template
from src.schemas.claim import Claim, SemanticReview
from src.schemas.evidence import Evidence, EvidenceLocation, ExtractionMethod, ReadingDepth
from src.schemas.outline import Outline, OutlineSection
from src.schemas.project import Project, ProjectArtifact
from src.schemas.review import ReviewItem, ReviewQueue
from src.schemas.source import Source, SourceState
from src.tools.docx_generator import DocxGenerationRequest, DocxGenerationTool
from src.tools.quantitative_review_workbook import build_quantitative_review_workbook, source_to_review_record
from src.tools.source_content import inspect_source, recheck_quote
from src.workflows.academic import AcademicWritingRequest, AcademicWritingResponse, AcademicWritingWorkflow
from src.workflows.export_bundle import export_bundle
from src.workflows.gates import check_academic_integrity, check_document_quality


def project_at(root: Path):
    path = root / "audit_project"
    path.mkdir(parents=True, exist_ok=True)
    return Project(name="audit_project", workspace="test", path=str(path), title="Synthetic test")


def test_search_counts_come_from_stage_records_not_claimed_totals(tmp_path):
    records, template = assessed_records(1, tmp_path, _template())
    log = tmp_path / "search_log.json"
    rows = [{"id": f"src{i}", "title": f"Synthetic candidate {i}"} for i in range(3)]
    payload = {"recorded_at": "2026-10-09", "origin": "synthetic search test", "records_by_stage": {
        "found": rows + [rows[0]], "deduplicated": rows, "screened": rows[:2]}}
    log.write_text(json.dumps(payload), encoding="utf-8")
    summary = {"found": 999, "search_log_path": str(log), "search_log_sha256": hashlib.sha256(log.read_bytes()).hexdigest()}
    result = build_quantitative_review_workbook(records, tmp_path / "review.xlsx", template_path=template, summary=summary)
    assert [result.counts[k] for k in ("found", "deduplicated", "screened")] == [4, 3, 2]
    assert {"field": "found", "input": 999, "actual": 4} in result.counts["reconciliation_findings"]
    payload.pop("records_by_stage"); payload["found"] = 999
    log.write_text(json.dumps(payload), encoding="utf-8")
    summary["search_log_sha256"] = hashlib.sha256(log.read_bytes()).hexdigest()
    result = build_quantitative_review_workbook(records, tmp_path / "untraceable.xlsx", template_path=template, summary=summary)
    assert result.counts["found"] == 1 and any(f["field"] == "search_log" for f in result.counts["reconciliation_findings"])


def test_run_source_snapshots_and_export_survive_original_file_change(tmp_path):
    project = project_at(tmp_path); sources, claims, ev, outline = bundle(project.directory)
    result = AcademicWritingWorkflow().execute(AcademicWritingRequest(project=project, sources=sources, claims=claims, evidence=ev, outline=outline))
    assert result.success
    snapshot = Source.model_validate(json.loads((Path(result.run_dir) / "sources_snapshot.json").read_text(encoding="utf-8"))[0])
    assert snapshot.retrieval_path != sources[0].retrieval_path and inspect_source(snapshot)["full_text"]
    original = Path(sources[0].retrieval_path); original.write_text("Changed after the run", encoding="utf-8")
    Path(sources[0].metadata["verification_artifact"]["path"]).write_text("Changed metadata", encoding="utf-8")
    from src.tools.source_content import stored_metadata
    assert inspect_source(snapshot)["full_text"] and stored_metadata(snapshot)
    exported = export_bundle(project.directory, run_id=result.run_id)
    assert exported["success"]
    with zipfile.ZipFile(exported["bundle_path"]) as archive:
        assert archive.testzip() is None and any(name.startswith("sources/") for name in archive.namelist())


def test_late_failure_retains_partial_workbook_and_reports_failed_runtime(tmp_path, monkeypatch):
    project = project_at(tmp_path); sources, claims, ev, outline = bundle(project.directory)
    import src.workflows.academic as module
    monkeypatch.setattr(module, "check_document_quality", lambda *args: (_ for _ in ()).throw(ValueError("synthetic quality failure")))
    result = AcademicWritingWorkflow().execute(AcademicWritingRequest(project=project, sources=sources, claims=claims, evidence=ev,
        outline=outline, quantitative_review=True, quantitative_review_template=str(_template())))
    assert not result.success and result.metadata["result_status"] == "FAILED" and not result.metadata["finalization_allowed"]
    assert Path(result.quantitative_workbook_path).is_file() and result.draft_path
    report = Path(result.quantitative_report_path).read_text(encoding="utf-8")
    assert "Status: FAILED" in report and "Status runtime: failed" in report
    summary = json.loads((Path(result.run_dir) / "run_summary.json").read_text(encoding="utf-8"))
    assert any(a["name"] == "quantitative_review.xlsx" for a in summary["artifacts"])


def bundle(root: Path, count=1):
    sources, claims, evidence = [], [], []
    for i in range(count):
        text = f"Synthetic observation for group {chr(65+i)}."
        source = fixture_source(Source(id=f"src_audit_{i}", title=f"Synthetic study {i}", authors=[f"Fixture{chr(65+i)}, A."],
            year=2025, venue="Synthetic test journal", state=SourceState.APPROVED), text, root / "source_documents", cache_report=True)
        claim = Claim(id=f"clm_audit_{i}", claim_text=text, status="SUPPORTED", supporting_sources=[source.id], supporting_evidence=[f"evd_audit_{i}"])
        ev = Evidence(id=f"evd_audit_{i}", claim_id=claim.id, source_id=source.id, evidence_text=text,
            extraction_method=ExtractionMethod.VERBATIM_ABSTRACT, location=EvidenceLocation(locator="abstract"), quote_verified=True)
        sources.append(source); claims.append(claim); evidence.append(ev)
    outline = Outline(title="Synthetic test", sections=[OutlineSection(title="Observations", claim_ids=[c.id for c in claims])])
    return sources, claims, evidence, outline


def test_stale_template_links_and_comments_are_cleared(tmp_path):
    w = load_workbook(_template()); m = w.worksheets[0]
    for col in (14, 15, 16):
        m.cell(6, col).hyperlink = "https://example.test/old-article"
        m.cell(6, col).comment = Comment("old example", "example")
        m.cell(7, col).hyperlink = "https://example.test/old-article"
    template = tmp_path / "template.xlsx"; w.save(template)
    result = build_quantitative_review_workbook([{"title": "Current article", "article_url": "https://example.test/current"},
        {"title": "No URL"}], tmp_path / "review.xlsx", template_path=template)
    reopened = load_workbook(result.workbook_path); m = reopened.worksheets[0]
    assert m.cell(6, 15).hyperlink.target == m.cell(6, 15).value
    for col in (14, 16):
        assert m.cell(6, col).hyperlink is None and m.cell(6, col).comment is None
    for col in (14, 15, 16):
        assert m.cell(7, col).hyperlink is None


def test_full_text_label_needs_retrieval_and_examination(tmp_path):
    sources, claims, ev, _ = bundle(tmp_path)
    source = sources[0]; source.reading_depth = ReadingDepth.FULL_TEXT
    source.retrieval_path = str(tmp_path / "missing.pdf")
    result = check_academic_integrity(claims=claims, evidence=ev, sources=sources)
    assert not result.ok and any("FULL_TEXT" in finding for finding in result.review_reasons)
    record = source_to_review_record(source)
    record.update(full_text_verified=True, provenance={"random": "assertion"}, access_status="full-text berhasil dibaca")
    workbook = build_quantitative_review_workbook([record], tmp_path / "review.xlsx", template_path=_template())
    assert workbook.counts["full_text_read"] == 0
    assert load_workbook(workbook.workbook_path).worksheets[0].cell(6, 17).value == "perlu review manusia"


def test_file_exists_does_not_prove_full_reading(tmp_path):
    source = bundle(tmp_path)[0][0]
    proof = inspect_source(source)
    assert proof["full_text"] and not proof["fully_read"]
    source.metadata["examination"] = {"reviewer": "claimed reviewer", "scope": "full"}
    assert not inspect_source(source)["fully_read"]


def test_false_quote_flag_is_recomputed_and_body_location_wins(tmp_path):
    sources, _, ev, _ = bundle(tmp_path)
    evidence = ev[0]; evidence.evidence_text = "This invented quote never occurs."
    assert not recheck_quote(sources[0], evidence) and not evidence.quote_verified
    evidence.evidence_text = "Fixture result."
    evidence.extraction_method = ExtractionMethod.VERBATIM_FULLTEXT
    evidence.location = EvidenceLocation(section="Results")
    assert recheck_quote(sources[0], evidence)
    evidence.location = EvidenceLocation(section="Methods")
    assert not recheck_quote(sources[0], evidence)


def test_fully_examined_source_can_use_abstract_evidence(tmp_path):
    records, _ = assessed_records(1, tmp_path, _template())
    source = Source.model_validate(records[0]["source_snapshot"])
    source.reading_depth = ReadingDepth.FULL_TEXT
    claim = Claim(id="clm", claim_text=source.abstract, status="SUPPORTED", supporting_sources=[source.id], supporting_evidence=["evd"])
    ev = Evidence(id="evd", claim_id=claim.id, source_id=source.id, evidence_text=source.abstract,
        extraction_method="VERBATIM_ABSTRACT", reading_depth="ABSTRACT_ONLY", location=EvidenceLocation(locator="abstract"))
    assert inspect_source(source)["fully_read"]
    assert check_academic_integrity(claims=[claim], evidence=[ev], sources=[source]).ok
    assert source.abstract == "Synthetic methods evidence."


@pytest.mark.parametrize("output_type", ["extraction_matrix", "search_report", "background"])
def test_document_topology_is_scoped(tmp_path, output_type):
    project = project_at(tmp_path); project.output_type = output_type
    assert check_document_quality("# Output\n## Findings\nA source-grounded observation.", project)["passed"]
    project.output_type = "literature_review"
    result = DocxGenerationTool().execute(DocxGenerationRequest(project=project,
        draft="# Review\n## Pendahuluan\nObservation.\n## Metode\n\n## Kesimpulan\n",
        citation_audit_passed=True, fact_audit_passed=True))
    assert not result.success and not (project.directory / "final.docx").exists()
    assert "metode" in result.error_message.casefold()


def test_ten_available_seven_cited_and_docx_references_agree(tmp_path):
    from docx import Document
    project = project_at(tmp_path)
    sources, claims, ev, outline = bundle(project.directory, 10)
    outline.sections[0].claim_ids = [c.id for c in claims[:7]]
    response = AcademicWritingWorkflow().execute(AcademicWritingRequest(project=project, sources=sources,
        claims=claims, evidence=ev, outline=outline))
    assert response.success and response.metadata["finalization_allowed"]
    refs = response.metadata["references"]
    assert refs["available"] == 10 and refs["cited"] == 7 and len(refs["written"]) == 7
    paragraphs = [p.text for p in Document(response.docx_path).paragraphs]
    for ref in refs["written"]:
        assert paragraphs.count(ref["formatted"]) == 1
    assert "FixtureH" not in "\n".join(paragraphs)


def test_partial_workbook_propagates_without_claiming_technical_failure(tmp_path):
    project = project_at(tmp_path); sources, claims, ev, outline = bundle(project.directory)
    response = AcademicWritingWorkflow().execute(AcademicWritingRequest(project=project, sources=sources,
        claims=claims, evidence=ev, outline=outline, quantitative_review=True,
        quantitative_review_template=str(_template()), quantitative_review_records=[source_to_review_record(sources[0])]))
    assert response.success and response.metadata["execution_success"]
    assert response.metadata["result_status"] == "PARTIAL" and not response.metadata["finalization_allowed"]
    assert response.needs_human_review and response.docx_path is None
    report = Path(response.quantitative_report_path).read_text(encoding="utf-8")
    assert '"finalization_allowed": false' in report and "100%" not in report
    summary = json.loads((Path(response.run_dir) / "run_summary.json").read_text(encoding="utf-8"))
    assert summary["result_status"] == "PARTIAL" and summary["needs_human_review"]


def test_summary_does_not_overwrite_actual_75_records(tmp_path):
    records, template = assessed_records(75, tmp_path, _template())
    result = build_quantitative_review_workbook(records, tmp_path / "review.xlsx", template_path=template, summary={"eligible": 999})
    assert result.counts["eligible"] == 75 and result.ranked_count == 75
    assert result.counts["reconciliation_findings"] == [{"field": "eligible", "input": 999, "actual": 75}]
    assert result.status == "PARTIAL"


def test_extension_preserves_formula_ranges_and_style(tmp_path):
    records, template = assessed_records(75, tmp_path, _template())
    w = load_workbook(template); m = w.worksheets[0]
    m["R6"] = "=A6*2"; m["R7"] = "=A7*2"
    m.auto_filter.ref = "A5:Q45"
    m.print_area = "A1:R48"
    m.add_table(Table(displayName="ArticleData", ref="A5:Q45"))
    validation = DataValidation(type="whole", operator="greaterThan", formula1="0"); validation.add("A6:A45"); m.add_data_validation(validation)
    w.defined_names.add(DefinedName("ArticleRows", attr_text="'Master 40 Kuantitatif PoPCites'!$A$6:$Q$45"))
    w.save(template)
    result = build_quantitative_review_workbook(records, tmp_path / "review.xlsx", template_path=template)
    reopened = load_workbook(result.workbook_path); m = reopened.worksheets[0]
    assert m["R46"].value == "=A46*2" and m["R80"].value == "=A80*2"
    assert m.auto_filter.ref == "A5:Q80" and m.tables["ArticleData"].ref == "A5:Q80"
    assert str(m.data_validations.dataValidation[0].sqref) == "A6:A80"
    assert "83" in m.print_area and "Q80" in reopened.defined_names["ArticleRows"].attr_text


def test_explicit_selection_keeps_corpus_and_excludes_ineligible(tmp_path):
    records, template = assessed_records(41, tmp_path, _template())
    records[8]["eligible"] = False  # High-scoring article must still be excluded.
    result = build_quantitative_review_workbook(records, tmp_path / "review.xlsx", template_path=template, request_text="tepat 40 artikel")
    assert result.article_count == 41 and result.ranked_count == 40
    ranking = load_workbook(result.workbook_path).worksheets[2]
    assert 9 not in [ranking.cell(5+i, 2).value for i in range(40)]
    screening = json.loads((tmp_path / "screening_audit.json").read_text(encoding="utf-8"))
    assert len(screening) == 41 and screening[8]["reasons"]
    records[7]["eligible"] = False
    deficit = build_quantitative_review_workbook(records, tmp_path / "deficit.xlsx", template_path=template, exact_count=40)
    assert deficit.ranked_count == 39 and deficit.counts["selection_deficit"] == 1 and deficit.status == "PARTIAL"


def test_high_pending_blocks_final_but_resolved_and_warnings_do_not(tmp_path):
    project = project_at(tmp_path); sources, claims, ev, outline = bundle(project.directory)
    queue = ReviewQueue([ReviewItem(item_type="source", item_id=sources[0].id, severity="HIGH", reason="Evidence requires review", recommended_action="Inspect snapshot")])
    queue.save(project.artifact_path(ProjectArtifact.REVIEW_QUEUE), root=project.directory)
    request = AcademicWritingRequest(project=project, sources=sources, claims=claims, evidence=ev, outline=outline)
    blocked = AcademicWritingWorkflow().execute(request)
    assert blocked.success and not blocked.metadata["finalization_allowed"] and blocked.docx_path is None
    queue.items[0].status = "RESOLVED"
    queue.add(ReviewItem(item_type="system", item_id="warning", severity="LOW", reason="Optional note", recommended_action="Read note"))
    queue.save(project.artifact_path(ProjectArtifact.REVIEW_QUEUE), root=project.directory)
    resolved = AcademicWritingWorkflow().execute(request)
    assert resolved.metadata["finalization_allowed"] and resolved.docx_path


def test_export_uses_requested_run_and_manifest_matches_actual_files(tmp_path):
    project = project_at(tmp_path); sources, claims, ev, outline = bundle(project.directory)
    request = AcademicWritingRequest(project=project, sources=sources, claims=claims, evidence=ev, outline=outline)
    first = AcademicWritingWorkflow().execute(request)
    immutable = Path(first.run_dir) / "artifacts" / "draft.md"
    original = immutable.read_bytes()
    request.outline.title = "Changed later draft"
    second = AcademicWritingWorkflow().execute(request)
    assert second.run_id != first.run_id
    exported = export_bundle(project.directory, run_id=first.run_id)
    assert exported["success"]
    with zipfile.ZipFile(exported["bundle_path"]) as z:
        assert z.testzip() is None and z.read("academic_output/draft.md") == original
    manifest = json.loads(Path(exported["manifest_path"]).read_text(encoding="utf-8"))
    for artifact in manifest["artifacts"] + [manifest["zip"]]:
        path = Path(artifact["path"])
        assert hashlib.sha256(path.read_bytes()).hexdigest() == artifact["sha256"]
        assert path.stat().st_size == artifact["size"]


def test_cli_api_finalize_forward_same_review_options(tmp_path, monkeypatch):
    from src.runtime import cli, monitor
    project = project_at(tmp_path)
    options = {"quantitative_review": True, "quantitative_review_template": str(_template()),
        "quantitative_review_records": [{"title": "Synthetic candidate"}], "quantitative_review_best_count": 10,
        "require_free_full_text": True}
    project.research_options = options
    (project.directory / "project.json").write_text(project.model_dump_json(), encoding="utf-8")
    (project.directory / "workflow_options.json").write_text(json.dumps(options), encoding="utf-8")
    payload = {"project": project.model_dump(mode="json"), **options}
    input_path = tmp_path / "input.json"; input_path.write_text(json.dumps(payload), encoding="utf-8")
    observed = []
    def capture(self, request):
        observed.append(request)
        return AcademicWritingResponse(success=True)
    monkeypatch.setattr(AcademicWritingWorkflow, "execute", capture)
    monkeypatch.setattr(cli, "_project_from_payload", lambda *a, **kw: project)
    monkeypatch.setattr(monitor, "_project_from_payload", lambda *a, **kw: project)
    cli._cmd_run_academic(argparse.Namespace(input_json=str(input_path), no_docx=True, resume=False))
    cli._cmd_finalize(argparse.Namespace(project=str(project.directory), project_flag=None, input_json=None, no_docx=True))
    handler = object.__new__(monitor.MonitorHandler)
    raw = json.dumps(payload).encode()
    handler.path = "/api/run-academic"; handler.headers = {"Content-Length": str(len(raw))}
    handler.rfile = io.BytesIO(raw)
    handler._token_ok = lambda: True; handler._json = lambda *args: None
    handler.do_POST()
    assert len(observed) == 3
    for request in observed:
        for name, value in options.items():
            assert getattr(request, name) == value


def test_resume_retains_project_configuration_and_semantic_reviews(tmp_path, monkeypatch):
    from src.runtime import cli
    from src.core.paths import SystemPaths
    from src.core.config import SystemConfig
    project = project_at(tmp_path); project.title = "Retained title"; project.language = "en"
    project.workspace = "TUGAS 1"
    project.research_options = {"quantitative_review": False}
    (project.directory / "project.json").write_text(project.model_dump_json(), encoding="utf-8")
    paths = SystemPaths(workspace_root=tmp_path, system_root=tmp_path / "DATA BASE")
    monkeypatch.setattr(cli, "get_paths", lambda: paths)
    monkeypatch.setattr(cli, "get_config", lambda: SystemConfig())
    resumed = cli._project_from_payload({"project": project.model_dump(mode="json"), "resume": True}, command="run-academic")
    assert resumed.directory == project.directory and resumed.language == "en" and resumed.title == "Retained title"
    sources, claims, ev, outline = bundle(project.directory)
    review = SemanticReview(claim_id="clm_existing_review", decision="SUPPORTED", reason="Existing synthetic review")
    (project.directory / "semantic_reviews.json").write_text(json.dumps([review.model_dump(mode="json")]), encoding="utf-8")
    request = AcademicWritingRequest(project=resumed, sources=sources, claims=claims, evidence=ev, outline=outline, generate_docx=False)
    first = AcademicWritingWorkflow().execute(request); second = AcademicWritingWorkflow().execute(request)
    assert first.run_id != second.run_id and first.success and second.success
    persisted = json.loads((project.directory / "semantic_reviews.json").read_text(encoding="utf-8"))
    assert persisted[0]["claim_id"] == review.claim_id and persisted[0]["reason"] == review.reason


@pytest.mark.parametrize("kind", ["html", "pdf", "json"])
def test_artifacts_are_parsed_with_their_actual_format(tmp_path, kind):
    source = bundle(tmp_path)[0][0]
    text = source.title + "\nMethods\nA synthetic method.\nResults\nA synthetic result.\nDiscussion\nA synthetic discussion.\nReferences\nSynthetic references."
    path = tmp_path / ("article." + kind)
    if kind == "html":
        path.write_text("<article>" + "".join(f"<p>{line}</p>" for line in text.splitlines()) + "</article>", encoding="utf-8")
    elif kind == "json":
        path.write_text(json.dumps({"full_text": text, "pages": [{"page": 1, "text": text}]}), encoding="utf-8")
    else:
        from pypdf import PdfWriter
        from pypdf.generic import NameObject, DictionaryObject, DecodedStreamObject
        writer = PdfWriter(); page = writer.add_blank_page(500, 500)
        font = DictionaryObject({NameObject("/Type"): NameObject("/Font"), NameObject("/Subtype"): NameObject("/Type1"), NameObject("/BaseFont"): NameObject("/Helvetica")})
        page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})})
        stream = DecodedStreamObject()
        stream.set_data(("BT /F1 12 Tf 20 480 Td " + " ".join(f"({line}) Tj 0 -18 Td" for line in text.splitlines()) + " ET").encode())
        page[NameObject("/Contents")] = stream
        with path.open("wb") as out: writer.write(out)
    source.retrieval_path = str(path); source.metadata["retrieval"]["sha256"] = hashlib.sha256(path.read_bytes()).hexdigest()
    proof = inspect_source(source)
    assert proof["full_text"] and proof["readable"] and not proof["fully_read"]
    ev = Evidence(claim_id="clm", source_id=source.id, evidence_text="A synthetic result.",
        extraction_method="VERBATIM_FULLTEXT", location=EvidenceLocation(section="Results"))
    assert recheck_quote(source, ev)


def test_preview_and_abstract_do_not_become_full_manuscripts(tmp_path):
    source = bundle(tmp_path)[0][0]
    path = Path(source.retrieval_path)
    path.write_text(source.title + "\nAbstract\nThe methods, results, discussion and references are summarized here.", encoding="utf-8")
    source.metadata["retrieval"]["sha256"] = hashlib.sha256(path.read_bytes()).hexdigest()
    assert not inspect_source(source)["full_text"]
    path.write_text(source.title + "\nMethods\nResults\nReferences\nPreview only. Please log in.", encoding="utf-8")
    source.metadata["retrieval"]["sha256"] = hashlib.sha256(path.read_bytes()).hexdigest()
    assert not inspect_source(source)["full_text"]


def test_cached_verification_is_reused_only_while_hash_and_identity_match(tmp_path):
    from src.tools.verification_tool import VerificationEngine
    source = bundle(tmp_path)[0][0]
    class NoLookup:
        name = "unavailable provider"
        def lookup_by_bibliographic(self, **kw):
            raise AssertionError("A valid cached snapshot should avoid a new search")
    engine = VerificationEngine(providers=[NoLookup()])
    assert engine.verify(source).recommended_state == "METADATA_VERIFIED"
    source.metadata["verification_artifact"]["sha256"] = "invalid"
    assert engine.verify(source).recommended_state == "NEEDS_HUMAN_REVIEW"


def test_access_policy_filters_science_and_access_independently(tmp_path):
    project = project_at(tmp_path)
    records, template = assessed_records(3, project.directory, _template())
    sources = [Source.model_validate(r["source_snapshot"]) for r in records]
    sources[1].access_mode = "BORROW_ONLY"
    sources[2].metadata.pop("eligibility_review")
    response = AcademicWritingWorkflow().execute(AcademicWritingRequest(project=project, sources=sources,
        generate_docx=False, require_free_full_text=True))
    screen = json.loads((Path(response.run_dir) / "access_screening.json").read_text(encoding="utf-8"))
    assert [r["included"] for r in screen] == [True, False, False]
    assert screen[1]["proof"]["scientific_eligible"] and not screen[1]["proof"]["legal_free"]
    assert screen[2]["proof"]["legal_free"] and not screen[2]["proof"]["scientific_eligible"]
    assert not response.metadata["finalization_allowed"]
    assert len(json.loads((Path(response.run_dir) / "sources_snapshot.json").read_text(encoding="utf-8"))) == 3


def test_export_with_old_input_resolves_new_execution_project(tmp_path, monkeypatch):
    from src.runtime import cli
    from src.core.paths import SystemPaths
    from src.core.config import SystemConfig
    origin = tmp_path / "original"; origin.mkdir()
    actual = tmp_path / "TUGAS 1" / "new_task"; (actual / "runs" / "run_actual").mkdir(parents=True)
    project = Project(name="new_task", workspace="TUGAS 1", path=str(actual), origin_project_path=str(origin))
    (actual / "project.json").write_text(project.model_dump_json(), encoding="utf-8")
    payload = tmp_path / "old_input.json"; payload.write_text(json.dumps({"project_path": str(origin)}), encoding="utf-8")
    monkeypatch.setattr(cli, "get_paths", lambda: SystemPaths(workspace_root=tmp_path, system_root=tmp_path / "DATA BASE"))
    monkeypatch.setattr(cli, "get_config", lambda: SystemConfig())
    assert cli._resolve_project_dir(argparse.Namespace(input_json=str(payload), run_id="run_actual")) == actual


def test_manifest_write_failure_is_reported(tmp_path, monkeypatch):
    project = project_at(tmp_path); sources, claims, ev, outline = bundle(project.directory)
    response = AcademicWritingWorkflow().execute(AcademicWritingRequest(project=project, sources=sources, claims=claims, evidence=ev, outline=outline))
    real_write = Path.write_text
    def fail_manifest(path, *args, **kw):
        if path.name.endswith(".manifest.json"):
            raise OSError("Synthetic write failure")
        return real_write(path, *args, **kw)
    monkeypatch.setattr(Path, "write_text", fail_manifest)
    result = export_bundle(project.directory, run_id=response.run_id)
    assert not result["success"] and "manifest write failed" in result["error"].casefold()


def test_invalid_template_stops_before_manuscript_execution(tmp_path):
    project = project_at(tmp_path); sources, claims, ev, outline = bundle(project.directory)
    template = tmp_path / "invalid.xlsx"; template.write_bytes(b"not an xlsx")
    response = AcademicWritingWorkflow().execute(AcademicWritingRequest(project=project, sources=sources, claims=claims,
        evidence=ev, outline=outline, quantitative_review=True, quantitative_review_template=str(template)))
    assert not response.success and response.metadata["execution_success"] is False
    assert not (project.directory / "draft.md").exists()


def test_direct_docx_cannot_bypass_partial_quantitative_workbook(tmp_path):
    project = project_at(tmp_path); project.research_options = {"quantitative_review": True}
    response = DocxGenerationTool().execute(DocxGenerationRequest(project=project, draft="# Test\n## Findings\nActual content.",
        citation_audit_passed=True, fact_audit_passed=True, quantitative_workbook_status="PARTIAL"))
    assert not response.success and not (project.directory / "final.docx").exists()


def test_direct_docx_rechecks_source_proof_and_project_access_policy(tmp_path):
    from src.tools.reference_formatter import format_reference_list
    project = project_at(tmp_path); source = bundle(tmp_path)[0][0]
    refs = format_reference_list([source])
    draft = f"# Test\n## Findings\nA source-grounded observation (FixtureA, 2025)."
    request = DocxGenerationRequest(project=project, draft=draft, sources=[source], reference_list=refs,
        citation_audit_passed=True, fact_audit_passed=True)
    source.reading_depth = "FULL_TEXT"; source.retrieval_path = str(tmp_path / "missing.pdf")
    assert not DocxGenerationTool().execute(request).success
    source.reading_depth = "ABSTRACT_ONLY"
    project.research_options = {"require_free_full_text": True}
    assert not DocxGenerationTool().execute(request).success
    source.metadata.pop("verification_artifact")
    project.research_options = {}
    assert not DocxGenerationTool().execute(request).success


def test_deep_research_reports_the_actual_post_audit_review_queue(tmp_path):
    from src.workflows.deep_research import DeepResearchRequest, DeepResearchWorkflow
    project = project_at(tmp_path); sources, claims, ev, outline = bundle(project.directory)
    ev[0].evidence_text = "An invented quotation absent from every snapshot."
    response = DeepResearchWorkflow().execute(DeepResearchRequest(project=project, user_request="Synthetic audit queue check",
        sources=sources, claims=claims, evidence=ev, outline=outline, generate_docx=False))
    actual = ReviewQueue.load(project.artifact_path(ProjectArtifact.REVIEW_QUEUE))
    assert not response.metadata["finalization_allowed"] and actual.blocking_items()
    assert {i.reason for i in actual.blocking_items()} == {i["reason"] for i in response.metadata["blocking_review_items"]}
    stored = ReviewQueue.load(Path(response.run_dir) / "review_queue.json")
    assert {i.reason for i in stored.blocking_items()} == {i.reason for i in actual.blocking_items()}
    assert (Path(response.run_dir) / "search_log.json").is_file()


def test_failed_provider_is_audited_and_never_reported_as_completed_research(tmp_path):
    from src.agents.research import DiscoveryAgent, DiscoveryRequest
    from src.tools.research_tool import ResearchResponse
    from src.workflows.deep_research import DeepResearchRequest, DeepResearchWorkflow
    class DeniedProvider:
        name = "synthetic denied provider"
        def execute(self, request):
            return ResearchResponse(success=False, error_code="HTTP_403", error_message="Synthetic access denied", http_status=403)
    discovery = DiscoveryAgent().execute(DiscoveryRequest(queries=["synthetic query"], providers=[DeniedProvider()]))
    assert not discovery.success and discovery.metadata["provider_failures"] and discovery.needs_human_review
    project = project_at(tmp_path)
    result = DeepResearchWorkflow().execute(DeepResearchRequest(project=project, user_request="Synthetic failed provider",
        providers=[DeniedProvider()], generate_docx=False, quantitative_review=True, quantitative_review_template=str(_template())))
    assert not result.success and result.metadata["execution_success"] is False and not result.metadata["finalization_allowed"]
    assert any("HTTP" in str(item) or "access denied" in item["reason"] for item in result.metadata["blocking_review_items"])
    assert "Status runtime: failed" in Path(result.metadata["quantitative_review"]["report_path"]).read_text(encoding="utf-8")
    log = json.loads((Path(result.run_dir) / "search_log.json").read_text(encoding="utf-8"))
    assert log["provider_failures"] and log["records_by_stage"]["found"] == []
